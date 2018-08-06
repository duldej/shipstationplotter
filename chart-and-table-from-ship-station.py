import os
import pandas as pd
import csv
import openpyxl
import xlsxwriter

#from datetime import datetime
#import matplotlib.pyplot as pyplot
import numpy as np
#import array
arr=np.array([])
#set working directory
os.chdir("/Users/Jason/Desktop/")
table=0
tablet1=0
n=0
dates=[]
items=[]
quantity=[]
table=[]

#get filename input
n=input("please enter top n results:")
x=input("Please enter csv filename:")
n=int(n)
#create a list
#with open(x) as csv_file:
df=pd.read_csv(x)
    
    #csv.reader(csv_file, delimiter=",")
 
#convert dates to usable date code

pd.to_datetime(df['OrderDate']).apply(lambda x: x.date())
df['Month'] = pd.DatetimeIndex(df['OrderDate']).month
df2=df[['Month','Description','Quantity']].copy()
maxmonths=max(df['Month'])
#drop rows with missing data
df2.dropna(inplace=True)
    

#create pivot table
table=df2.pivot_table(index='Month', columns='Description',values='Quantity',aggfunc='sum')

#create totals
table.loc['sum'] = table.sum()


#transpose matrix
tablet=table.T

#take top n values
tablet1 = tablet.sort_values('sum',ascending = False).head(n)

#eliminate the blanks in the table
tablet1=tablet1.fillna(0)
#convert pivot table to df
flattened = pd.DataFrame(tablet1.to_records())

# Create a Pandas Excel writer using XlsxWriter as the engine.
sheet_name = 'Sheet1'
writer     = pd.ExcelWriter('pandas_chart_line.xlsx', engine='xlsxwriter')
flattened.to_excel(writer, sheet_name=sheet_name)

# Access the XlsxWriter workbook and worksheet objects from the dataframe.
workbook  = writer.book
worksheet = writer.sheets[sheet_name]

# Create a chart object.
chart = workbook.add_chart({'type': 'line'})

# Configure the series of the chart from the dataframe data.
for i in range(0,n):
    col = i + 1
    chart.add_series({
        'name':       ['Sheet1', i+1,1],
        'categories': ['Sheet1',0,2,0,maxmonths],
        'values':     ['Sheet1',i+1,2,i+1,maxmonths ],
    })

# Configure the chart axes.
chart.set_x_axis({'name': 'Months'})
chart.set_y_axis({'name': 'Units Sold', 'major_gridlines': {'visible': False}})

# Insert the chart into the worksheet.
worksheet.insert_chart('K2', chart)

# Close the Pandas Excel writer and output the Excel file.
writer.save()

















#write table to Excel file

#writer = pd.ExcelWriter('report.xlsx')

#tablet1.to_excel(writer, 'Sheet1')

#write chart to Excel file

#workbook=writer.book
#worksheet=writer.sheets['Sheet1']
#chart = workbook.add_chart({'type': 'line'})
#chart.add_series({'values':'=Sheet1!b2:g6'})
#worksheet.insert_chart('j1',chart)
#save and close the Excel file
#new approach to plotting the table
# Access the XlsxWriter workbook and worksheet objects from the dataframe.
workbook  = writer.book
worksheet = writer.sheets


#workbook = xlsxwriter.Workbook('chart.xlsx')
#worksheet = workbook.add_worksheet()


#worksheet.write_column('A1',tablet1[2])
#worksheet.write_column('b1',tablet1[3])
#worksheet.write_column('c1',tablet1[4])


# Create a chart object.
chart = workbook.add_chart({'type': 'line'})
descs=len(tablet1[1])
# Configure the series of the chart from the dataframe data.
#for i in range(0,descs):
#    col = i + 1
#    chart.add_series({
#        'name':       ['=Sheet1', 0, col],
#        'categories': ['=Sheet1', 0, 0,   descs, 0],
#        'values':     ['=Sheet1', 1, col, descs, col],
#    })
# Configure the chart. In simplest case we add one or more data series.
chart.add_series({'values': '=Sheet1!$b$2:$g$2'})
chart.add_series({'values': '=Sheet1!$B$3:$g$3'})
chart.add_series({'values': '=Sheet1!$b$4:$g$4'})






# Configure the chart axes.
chart.set_x_axis({'name': 'Index'})
chart.set_y_axis({'name': 'Value', 'major_gridlines': {'visible': False}})

# Insert the chart into the worksheet.
worksheet.insert_chart('L1', chart)



writer.save()
workbook.close()


#tablet2.groupby(['Month','Description']).count()['Quantity'].unstack().plot(kind='line')


#tablet2=pd.DataFrame(tablet)
#tablet2.plot()
#m=max(df2['Month'])
#for xlab in range(1,m):
 #   tablet1.plot(x=xlab, y='Description')
